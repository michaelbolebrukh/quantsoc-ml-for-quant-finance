"""Approximate PNG previews of deck/quantsoc-ml-workshop-2.pptx, for QA where LibreOffice cannot render.

Run: python3 deck/qa_render.py [out_dir]   (default out_dir: deck/qa_preview)
It re-reads the built .pptx with python-pptx and draws each shape with Pillow: the slide background, fills and
hairlines with their alpha (the deck's glass cards are a few percent of white over #0a0a0a), pictures, tables,
connectors, XY charts (series lines and markers on the stored axis range, in the stored colours), and text
wrapped with Inter from deck/assets/fonts, picking the weight from the run's font name ("Inter SemiBold") and
bold flag, so the overflow check measures the face the deck sets. It prints OVERFLOW for any text that needs
more height than its box, and OFFSLIDE for any shape past the slide edge. Exit code 1 if anything is flagged.
Blind spots: it is not PowerPoint. Chart axes, legends and labels are only sketched, autofit is ignored, dashes
are drawn solid, and glyphs missing from Inter are drawn by a fallback font. LibreOffice Impress, where it is
installed, is the better check: soffice --headless --convert-to pdf deck/quantsoc-ml-workshop-2.pptx.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu

DECK = Path(__file__).resolve().parent
PX = 120  # pixels per inch
EMU_IN = 914400
INTER = str(DECK / "assets" / "fonts" / "Inter-{}.ttf")
LIBERATION = "/usr/share/fonts/truetype/liberation/LiberationSans-{}.ttf"
SYMBOL = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
LINE = 1.21  # Inter's ascender + descender, in em (no line gap)
_cache = {}


def font(pt, weight):
    k = (round(pt, 1), weight)
    if k not in _cache:
        name = {"regular": "Regular", "semibold": "SemiBold", "bold": "Bold"}[weight]
        path = Path(INTER.format(name))
        if not path.exists():
            path = LIBERATION.format("Regular" if weight == "regular" else "Bold")
        _cache[k] = ImageFont.truetype(str(path), max(1, int(pt * PX / 72)))
    return _cache[k]


def sym_font(pt):
    p = Path(SYMBOL)
    return ImageFont.truetype(str(p), max(1, int(pt * PX / 72))) if p.exists() else font(pt, "bold")


def px(v):
    return int(round(v / EMU_IN * PX))


def _srgb(parent):
    """(#hex, alpha 0..255) from the first a:srgbClr under parent, or None."""
    if parent is None:
        return None
    el = parent.find(".//" + qn("a:srgbClr"))
    if el is None:
        return None
    a = el.find(qn("a:alpha"))
    alpha = int(int(a.get("val")) / 100000 * 255) if a is not None else 255
    return "#" + el.get("val"), alpha


def fill_of(shape):
    spPr = shape._element.find(qn("p:spPr"))
    if spPr is None:
        return None
    return _srgb(spPr.find(qn("a:solidFill")))


def line_of(shape):
    spPr = shape._element.find(qn("p:spPr"))
    ln = spPr.find(qn("a:ln")) if spPr is not None else None
    if ln is None or ln.find(qn("a:noFill")) is not None:
        return None, 0
    c = _srgb(ln.find(qn("a:solidFill")))
    w = int(ln.get("w", "12700"))
    return c, max(1, int(round(w / 12700 * PX / 72)))


def run_style(r):
    f = r.font
    pt = f.size.pt if f.size else 18
    weight = "bold" if f.bold else ("semibold" if (f.name or "").endswith("SemiBold") else "regular")
    try:
        c = "#" + str(f.color.rgb)
    except Exception:
        c = "#000000"
    return pt, weight, c


def layout_text(tf, width_px):
    """Return list of (lines, size, alignment, space_after); each line = list of (text, font, colour, space)."""
    out = []
    for p in tf.paragraphs:
        words = []   # each word: list of (text, font, colour) segments, so runs join without a space
        size = None
        glue = False
        for r in p.runs:
            pt, weight, c = run_style(r)
            size = size or pt
            fnt = font(pt, weight)
            pieces = r.text.split(" ")
            for k, w in enumerate(pieces):
                if k == 0 and glue and words and w:
                    words[-1].append((w, fnt, c))
                elif w:
                    words.append([(w, fnt, c)])
            glue = not r.text.endswith(" ")
        size = size or 18
        lines, cur, cur_w = [], [], 0
        for wsegs in words:
            fnt0 = wsegs[0][1]
            sp = fnt0.getlength(" ") if cur else 0
            ww = sum(f.getlength(t) for t, f, c in wsegs)
            if cur and cur_w + sp + ww > width_px:
                lines.append(cur); cur, cur_w = [], 0; sp = 0
            for j, (t, f, c) in enumerate(wsegs):
                cur.append((t, f, c, sp if j == 0 else 0))
            cur_w += sp + ww
        lines.append(cur)
        after = p.space_after.pt if p.space_after is not None else 0
        out.append((lines, size, p.alignment, after))
    return out


def draw_text(d, tf, x, y, w, h, label, flags):
    ml = px(tf.margin_left if tf.margin_left is not None else Emu(91440))
    mr = px(tf.margin_right if tf.margin_right is not None else Emu(91440))
    mt = px(tf.margin_top if tf.margin_top is not None else Emu(45720))
    mb = px(tf.margin_bottom if tf.margin_bottom is not None else Emu(45720))
    iw = w - ml - mr
    paras = layout_text(tf, iw)
    total = 0
    for lines, size, _, after in paras:
        total += len(lines) * size * LINE * PX / 72 + after * PX / 72
    total -= paras[-1][3] * PX / 72 if paras else 0
    ih = h - mt - mb
    if total > ih + 4:
        flags.append(f"OVERFLOW {label}: needs {total / PX:.2f} in, box {ih / PX:.2f} in")
    widest = max((sum(f.getlength(t) + sp for t, f, c, sp in ln) for lines, *_ in paras for ln in lines), default=0)
    if widest > iw + 4 and tf.word_wrap is not False:
        flags.append(f"OVERFLOW {label}: a word is {widest / PX:.2f} in wide, box {iw / PX:.2f} in")
    anchor = tf.vertical_anchor
    yy = y + mt
    if anchor == MSO_ANCHOR.MIDDLE:
        yy = y + mt + (ih - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        yy = y + h - mb - total
    for lines, size, align, after in paras:
        lh = size * LINE * PX / 72
        for ln in lines:
            lw = sum(f.getlength(t) + sp for t, f, c, sp in ln)
            xx = x + ml
            if align == PP_ALIGN.CENTER:
                xx = x + ml + (iw - lw) / 2
            elif align == PP_ALIGN.RIGHT:
                xx = x + w - mr - lw
            for t, f, c, sp in ln:
                xx += sp
                f2 = sym_font(f.size * 72 / PX) if any(ord(ch) > 0x2000 and ord(ch) not in (0x2192,) for ch in t) else f
                d.text((xx, yy + lh * 0.08), t, font=f2, fill=c)
                xx += f2.getlength(t)
            yy += lh
        yy += after * PX / 72


A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def draw_chart(d, gf, x, y, w, h):
    ch = gf.chart
    ns = {"c": C, "a": A}
    data = []
    for s in ch._chartSpace.findall(".//c:ser", ns):
        xs = [float(v.text) for v in s.findall("./c:xVal//c:v", ns)]
        ys = [float(v.text) for v in s.findall("./c:yVal//c:v", ns)]
        clr = s.find("./c:spPr/a:ln/a:solidFill/a:srgbClr", ns)
        mk = s.find(".//c:marker//a:srgbClr", ns)
        data.append((xs, ys, "#" + clr.get("val") if clr is not None else None, "#" + mk.get("val") if mk is not None else None))
    allx = [v for xs, *_ in data for v in xs] or [0, 1]
    ally = [v for _, ys, *_ in data for v in ys] or [0, 1]

    def rng(ax, vals):
        lo = ax.minimum_scale if ax.minimum_scale is not None else min(vals)
        hi = ax.maximum_scale if ax.maximum_scale is not None else max(vals)
        return lo, hi if hi > lo else lo + 1
    x0, x1 = rng(ch.category_axis, allx)
    y0, y1 = rng(ch.value_axis, ally)
    L, T, R, B = x + 60, y + 40, x + w - 20, y + h - 50
    axis = ch._chartSpace.find(".//c:valAx/c:spPr/a:ln/a:solidFill/a:srgbClr", ns)
    grid = ch._chartSpace.find(".//c:valAx/c:majorGridlines//a:srgbClr", ns)
    ac = "#" + axis.get("val") if axis is not None else "#9ca3af"
    if grid is not None:
        for k in range(1, 5):
            gy = B - (B - T) * k / 4
            d.line([(L, gy), (R, gy)], fill="#" + grid.get("val"), width=1)
    d.line([(L, B), (R, B)], fill=ac, width=1)
    d.line([(L, T), (L, B)], fill=ac, width=1)
    tc = "#acb6c3"
    f = font(10, "regular")
    d.text((L - 50, T), f"{y1:g}", font=f, fill=tc)
    d.text((L - 50, B - 14), f"{y0:g}", font=f, fill=tc)
    d.text((L, B + 6), f"{x0:g}", font=f, fill=tc)
    d.text((R - 30, B + 6), f"{x1:g}", font=f, fill=tc)
    for xs, ys, lc, mc in data:
        pts = [(L + (a - x0) / (x1 - x0) * (R - L), B - (b - y0) / (y1 - y0) * (B - T)) for a, b in zip(xs, ys)]
        if lc and len(pts) > 1:
            d.line(pts, fill=lc, width=3)
        if mc:
            for (a, b) in pts:
                d.ellipse([a - 3, b - 3, a + 3, b + 3], fill=mc)


def _rgba(c):
    if c is None:
        return None
    h, a = c
    return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5)) + (a,)


def render(pptx_path: Path, out_dir: Path) -> list[str]:
    prs = Presentation(str(pptx_path))
    W, H = px(prs.slide_width), px(prs.slide_height)
    out_dir.mkdir(parents=True, exist_ok=True)
    flags = []
    for n, slide in enumerate(prs.slides, 1):
        bg = "#ffffff"
        try:
            if slide.background.fill.type == 1:
                bg = "#" + str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        img = Image.new("RGBA", (W, H), bg)
        for i, sh in enumerate(slide.shapes):
            x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            label = f"slide {n} shape {i} ({sh.name})"
            if x < -2 or y < -2 or x + w > W + 2 or y + h > H + 2:
                flags.append(f"OFFSLIDE {label}")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA").resize((max(1, w), max(1, h)))
                img.alpha_composite(im, (x, y))
                continue
            layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
            d = ImageDraw.Draw(layer)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                draw_chart(d, sh, x, y, w, h)
            elif getattr(sh, "has_table", False) and sh.has_table:
                cy = y
                for r in sh.table.rows:
                    rh = px(r.height)
                    cx = x
                    for ci, c in enumerate(r.cells):
                        cw = px(sh.table.columns[ci].width)
                        fc = "#" + str(c.fill.fore_color.rgb) if c.fill.type == 1 else None
                        d.rectangle([cx, cy, cx + cw, cy + rh], fill=fc, outline="#3a3f49")
                        draw_text(d, c.text_frame, cx, cy, cw, rh, f"{label} cell", flags)
                        cx += cw
                    cy += rh
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fc = _rgba(fill_of(sh))
                lc, lw = line_of(sh)
                lc = _rgba(lc)
                kind = sh.auto_shape_type
                if kind == MSO_SHAPE.OVAL:
                    d.ellipse([x, y, x + w, y + h], fill=fc, outline=lc, width=lw if lc else 0)
                elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                    r = int(min(w, h) * (sh.adjustments[0] if len(sh.adjustments) else 0.16))
                    d.rounded_rectangle([x, y, x + w, y + h], r, fill=fc, outline=lc, width=lw if lc else 0)
                elif kind == MSO_SHAPE.RIGHT_ARROW:
                    d.polygon([(x, y + h * .3), (x + w * .6, y + h * .3), (x + w * .6, y), (x + w, y + h / 2),
                               (x + w * .6, y + h), (x + w * .6, y + h * .7), (x, y + h * .7)], fill=fc)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=fc, outline=lc, width=lw if lc else 0)
                if sh.has_text_frame and sh.text_frame.text.strip():
                    draw_text(d, sh.text_frame, x, y, w, h, label, flags)
            elif sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.__class__.__name__ == "Connector":
                ln = sh._element.find(".//" + qn("a:ln"))
                c = _rgba(_srgb(ln)) if ln is not None else None
                lw = max(1, int(round(int(ln.get("w", "12700")) / 12700 * PX / 72))) if ln is not None else 1
                d.line([(x, y), (x + w, y + h)], fill=c or (156, 163, 175, 255), width=lw)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                draw_text(d, sh.text_frame, x, y, w, h, label, flags)
            img.alpha_composite(layer)
        img.convert("RGB").save(out_dir / f"slide-{n:02d}.png")
    return flags


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DECK / "qa_preview"
    fl = render(DECK / "quantsoc-ml-workshop-2.pptx", out)
    print("\n".join(fl) if fl else "no overflow or off-slide shapes flagged")
    print(f"previews in {out}")
    sys.exit(1 if fl else 0)
