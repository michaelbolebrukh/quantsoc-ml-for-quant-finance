"""Build deck/quantsoc-ml-workshop-2.pptx and deck/SPEAKER_NOTES.md from deck/slides.yaml + deck/slides_bonus.yaml.

Run:  python3 deck/build_deck.py [--figures]      (from the build root or from deck/; paths are absolute)
Needs python-pptx, PyYAML, Pillow and segno. --figures (or a missing F1..F4 PNG) also needs matplotlib, the
quantsoc_ml package and data/prices.csv.gz, and takes about 30 seconds.

Content (words, notes, minute budgets, run markers, references, process stages) lives in slides.yaml (Part 1,
the main deck) and slides_bonus.yaml (the bonus section on learning theory, appended after the main deck's
last slide in the same file); this file is layout only. Every slide carries a stable `id`; slide numbers are
computed from position, and "{slide:<id>}" anywhere in the yaml renders as that slide's live number. The look is QuantSoc's own (next-pwa/product/DESIGN.md, next-pwa/app/styles/tokens.css):
dark-first canvas #0a0a0a, off-white --foreground, brand-blue accents, Inter, glass cards (a very low-alpha
foreground tint with a 1px hairline, since pptx cannot blur), radii that grow with elevation, semibold as the
emphasis weight. The colour constants below are the dark tokens; build() measures every text and indicator
pair against the WCAG floor and refuses to build below it.

Each build also regenerates the D_ concept diagrams (make_diagrams.py) and the QR codes (make_qr.py). The
result figures F1..F4 (deck/figures/F1_decile.png, F2_noise.png, F3_equity.png, F4_split.png) are drawn by
quantsoc_ml.notebook_support.make_deck_figures(dark=True); they are cached, and redrawn only when one is missing
or --figures is passed. If they cannot be drawn (no data or no matplotlib) a placeholder box names the file,
so the deck always builds. The same placeholder is drawn for any other missing result figure (the F5, F6 and
B1..B9 figures belong to deck/make_bonus_figures.py). The build fails loudly (non-zero exit) if a self-check
fails: slide counts per part, contiguous minute budgets per part (main 0 to MAIN_MINUTES, bonus 0 to
BONUS_MINUTES), run-marker and stage maps keyed by slide id, unknown {slide:id} references, an "In finance"
chip naming a row the comparison table does not have, a "Seen in Part 1" chip naming a slide outside Part 1,
footer labels, an em/en dash anywhere in slide text or notes, or a colour pair under its contrast floor.
"""
from __future__ import annotations

import gzip
import csv
import math
import random
import re
import sys
from pathlib import Path
from statistics import NormalDist

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DECK = Path(__file__).resolve().parent
ROOT = DECK.parent
sys.dont_write_bytecode = True
sys.path.insert(0, str(DECK))
import make_diagrams  # noqa: E402
import make_qr  # noqa: E402

OUT_PPTX = DECK / "quantsoc-ml-workshop-2.pptx"
OUT_NOTES = DECK / "SPEAKER_NOTES.md"
FIG = DECK / "figures"
ASSETS = DECK / "assets"
FONT_DIR = ASSETS / "fonts"          # Inter 4.1 (OFL, see Inter-LICENSE.txt): measuring here, rendering in QA
LOGO = ASSETS / "logo-dark.png"      # next-pwa/public/logos/logo-dark.png (the white mark), cropped to its ink
DATA = ROOT / "data" / "prices.csv.gz"
RESULT_FIGS = ("F1_decile.png", "F2_noise.png", "F3_equity.png", "F4_split.png")

# Self-check expectations, keyed by stable slide id (numbers move when slides are inserted).
MAIN_COUNT, BONUS_COUNT = 23, 22
MAIN_MINUTES, BONUS_MINUTES = 71, 25
FINANCE_TABLE = "textbook_vs_finance"   # the slide whose grid rows the "In finance" chips must name
EXPECTED_RUNS = {"data": 1, "target": 2, "leakage": 2, "baseline": 3, "time_split": 3, "forest": 4,
                 "overfit": 5, "portfolio": 6, "backtest": 7, "export": 8}
EXPECTED_STAGES = {
    # Part 1: the six process steps
    "title": "none", "map": "none", "textbook_vs_finance": "ahead", "data": "Prices", "data_card": "Prices",
    "target": "Features", "feature_universe": "Features", "literature": "Features", "features_drawn": "Features",
    "features_usable": "Features", "leakage": "Features", "baseline": "Model", "time_split": "Model",
    "tiny_edge": "Model", "tree": "Model", "forest": "Model", "deciles": "Model", "overfit": "Check",
    "trying_many": "Check", "portfolio": "Portfolio", "survivorship": "Backtest", "backtest": "Backtest",
    "export": "Backtest",
    # Bonus: the five chapters
    "bonus_divider": "ahead", "one_question": "Generalisation", "assumption": "Generalisation",
    "three_sets": "Honest testing", "cross_validation": "Honest testing", "test_once": "Honest testing",
    "bias_variance": "Trade-offs", "fit_three": "Trade-offs", "u_theory": "Trade-offs", "u_ours": "Trade-offs", "noise_floor": "Trade-offs",
    "regularisation": "Trade-offs", "double_descent": "Trade-offs", "family_map": "Model families",
    "properties": "Model families", "flex_interp": "Model families", "explaining": "Model families",
    "no_free_lunch": "Model families", "finance_rules": "Finance", "two_cultures": "Finance",
    "checklist": "Finance", "bonus_qr": "Finance",
}
PROCESS_STEPS = ["Prices", "Features", "Model", "Check", "Portfolio", "Backtest"]
BONUS_CHAPTERS = ["Generalisation", "Honest testing", "Trade-offs", "Model families", "Finance"]
SITE = "quant-soc.com"
WORDMARK = "QuantSoc"
SOCIETY = "Quantitative Finance Society"
REF_LABEL = "Reference"

# ---------------------------------------------------------------- QuantSoc dark tokens (tokens.css, .dark)
CANVAS = "0a0a0a"    # --background (neutral-950)
FG = "dae0e7"        # --foreground, 210 20% 88%
MUTED = "acb6c3"     # --muted-foreground, 215 16% 72%
DIM = "767f8c"       # future steps in the process strip: the dimmest text that still clears 4.5:1
RULE = "717a88"      # connectors, arrows, axis lines (non-text, 3:1 floor)
GRID = "232425"      # chart gridlines: --foreground at 12% over the canvas (decorative)
KEY = "61a6fa"       # --brand-400: labels, terms, step numbers
KEY_INK = "9ec3f0"   # key-text emphasis: --foreground warmed halfway to --brand-400 (see emphasis note)
PRIMARY = "6da2f8"   # dark --primary (217 91% 70%): filled pills, carrying CANVAS text as the site's buttons do
BRAND = "3b82f6"     # --brand-500, stock Tailwind blue-500: fills and marks only, never small text
WARN = "f26464"      # dark --destructive (0 84% 67%)
AMBER = "fccb4f"     # dark --warning (43 97% 65%), one chart series
WHITE = "ffffff"     # the QR tile only: a QR needs a light quiet zone
NONE_BAR = "3a3f49"  # "no position" bars in the portfolio picture (decorative)
GLASS, GLASS_LINE = 0.05, 0.14   # card tint and hairline alpha: the site's dark glass is 3% and 10%, raised a
                                 # little because a projector crushes near-black
FONT, FONT_SEMI = "Inter", "Inter SemiBold"
# Emphasis note: the site's key text is a vertical gradient from --foreground to --brand-400 per word. In pptx a
# text gradient spans the whole text frame, not the phrase, so a word on the first line of a paragraph would stay
# plain while one on the last line turned blue. A solid stop between the two, in semibold, is the honest emulation.

SW, SH = 13.333, 7.5
MX = 0.6                      # side margin
TOP_PLAIN, TOP_STRIP = 1.75, 1.98   # content top without and with the process strip
BOTTOM = 6.4                  # content bottom (reference footer below)
R_CARD, R_INPUT = 0.16, 0.08  # 16px cards, 8px inputs/tiles (slide width read as 1280px); pills are full-round


def rgb(h):
    return RGBColor.from_string(h)


def _hex(h):
    return [int(h[i:i + 2], 16) for i in (0, 2, 4)]


def mix(fg, bg, a):
    """fg at alpha a over bg, as the hex a viewer sees."""
    return "".join(f"{round(x * a + y * (1 - a)):02x}" for x, y in zip(_hex(fg), _hex(bg)))


def contrast(a, b):
    """WCAG 2.2 contrast ratio of two hex colours."""
    def lum(h):
        c = [v / 255 for v in _hex(h)]
        c = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    hi, lo = sorted((lum(a), lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


CARD = mix(FG, CANVAS, GLASS)           # what a glass card looks like over the canvas
PAST_CHIP = mix(FG, CANVAS, 0.06)
TILE = mix(BRAND, CARD, 0.20)
WARN_ROW = mix(WARN, CARD, 0.10)
CHIP = mix(BRAND, CANVAS, 0.20)         # the "In finance" / "Seen in Part 1" chip
HI_COL = mix(BRAND, CARD, 0.10)         # the highlighted column of a grid (the finance column on the table)
TEST_SET = mix(FG, CANVAS, 0.10)        # the test block in the train / validation / test picture
CONTRAST_PAIRS = [  # (what, foreground, background, floor)
    ("body text", FG, CANVAS, 4.5), ("body text on a glass card", FG, CARD, 4.5),
    ("muted text", MUTED, CANVAS, 4.5), ("muted text on a glass card", MUTED, CARD, 4.5),
    ("key-text emphasis", KEY_INK, CANVAS, 4.5), ("key-text emphasis on a card", KEY_INK, CARD, 4.5),
    ("brand-400 labels on a card", KEY, CARD, 4.5), ("number on a brand tile", KEY, TILE, 4.5),
    ("strip: current step text on primary", CANVAS, PRIMARY, 4.5), ("strip: past step", MUTED, PAST_CHIP, 4.5),
    ("strip: future step", DIM, CANVAS, 4.5), ("run pill text on primary", CANVAS, PRIMARY, 4.5),
    ("warning text", WARN, CANVAS, 4.5), ("warning row text", WARN, WARN_ROW, 4.5),
    ("placeholder tag text", CANVAS, WARN, 4.5), ("chart series amber", AMBER, CANVAS, 3.0),
    ("primary fill (current step, pill)", PRIMARY, CANVAS, 3.0), ("brand-500 fills (bars)", BRAND, CANVAS, 3.0),
    ("connectors, arrows, axes", RULE, CANVAS, 3.0), ("QR tile against the canvas", WHITE, CANVAS, 3.0),
    ("chip text", KEY_INK, CHIP, 4.5), ("chip label", KEY, CHIP, 4.5), ("chip hairline", KEY, CANVAS, 3.0),
    ("grid text on the highlighted column", FG, HI_COL, 4.5), ("grid header on the highlighted column", KEY, HI_COL, 4.5),
    ("test block text", FG, TEST_SET, 4.5), ("validation block text", KEY, TILE, 4.5),
    ("optional tag text", MUTED, CANVAS, 4.5),
]


# ---------------------------------------------------------------- low-level XML helpers
def _alpha(srgb_el, a):
    if a is not None and a < 1:
        for old in srgb_el.findall(qn("a:alpha")):
            srgb_el.remove(old)
        etree.SubElement(srgb_el, qn("a:alpha")).set("val", str(int(round(a * 100000))))


def set_fill(shape, colour, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(colour)
    _alpha(shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)


def set_line(line, colour, width_pt=0.75, alpha=None, dash=False):
    line.color.rgb = rgb(colour)
    line.width = Pt(width_pt)
    _alpha(line._get_or_add_ln().find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)
    if dash:
        line.dash_style = MSO_LINE_DASH_STYLE.DASH


def style_font(f, size, colour, weight="regular"):
    """weight: regular | semibold | bold. Semibold is the family 'Inter SemiBold', not a bold flag."""
    f.name = FONT_SEMI if weight == "semibold" else FONT
    f.size = Pt(size)
    f.bold = weight == "bold"
    f.color.rgb = rgb(colour)


# ---------------------------------------------------------------- text helpers
def _runs(text):
    """'a **b** c' -> [(a, False), (b, True), (c, False)]"""
    parts = text.split("**")
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


def plain(text):
    return text.replace("**", "")


def _fill_runs(p, text, size, colour, weight, emph_colour, tracking=None):
    for txt, emph in _runs(text):
        r = p.add_run()
        r.text = txt
        style_font(r.font, size, emph_colour if emph else colour,
                   "semibold" if (emph and weight == "regular") else weight)
        if tracking:
            r._r.get_or_add_rPr().set("spc", str(int(tracking)))


def add_text(slide, x, y, w, h, paras, size=20, color=FG, weight="regular", align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=8, emph_color=KEY_INK, line_spacing=None, tracking=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        _fill_runs(p, ptxt, size, color, weight, emph_color, tracking)
    return tb


_MEASURE = {}
_FALLBACK = {"regular": "LiberationSans-Regular.ttf", "semibold": "LiberationSans-Bold.ttf",
             "bold": "LiberationSans-Bold.ttf"}


def _mfont(weight):
    """Inter (the deck's face, shipped in deck/assets/fonts) at 200 px, for measuring wrapped text."""
    if weight not in _MEASURE:
        from PIL import ImageFont
        inter = FONT_DIR / {"regular": "Inter-Regular.ttf", "semibold": "Inter-SemiBold.ttf",
                            "bold": "Inter-Bold.ttf"}[weight]
        lib = Path("/usr/share/fonts/truetype/liberation") / _FALLBACK[weight]
        path = inter if inter.exists() else lib
        _MEASURE[weight] = ImageFont.truetype(str(path), 200) if path.exists() else None
    return _MEASURE[weight]


def est_lines(text, width_in, size_pt, weight="regular"):
    """Lines a paragraph needs at size_pt in width_in, measured with Inter's own metrics (emphasis in semibold)."""
    reg = _mfont(weight)
    emp = _mfont("semibold" if weight == "regular" else weight)
    if reg is None:
        cpl = max(8, int(width_in * 72 / (size_pt * 0.56)))
        return max(1, math.ceil(len(plain(text)) / cpl))
    limit = width_in * 72 / size_pt * 200 * 0.97   # px at 200 px per em-point unit
    words = []
    for seg, b in _runs(text):
        for k, wd in enumerate(seg.split(" ")):
            if k == 0 and words and not seg.startswith(" ") and not words[-1][1].endswith(" "):
                words[-1] = (words[-1][0] + (emp if b else reg).getlength(wd), words[-1][1] + wd)
            elif wd:
                words.append(((emp if b else reg).getlength(wd), wd))
    space = reg.getlength(" ")
    lines, cur = 1, 0.0
    for wl, _ in words:
        if cur and cur + space + wl > limit:
            lines, cur = lines + 1, wl
        else:
            cur += (space if cur else 0) + wl
    return lines


def text_h(n_lines, size):
    return n_lines * size * 1.22 / 72


def box(slide, x, y, w, h, fill=None, alpha=None, radius=R_CARD, line=None, line_alpha=None, line_w=0.75,
        dash=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    """radius in inches (a pill is radius >= h/2)."""
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    s.shadow.inherit = False
    if fill:
        set_fill(s, fill, alpha)
    else:
        s.fill.background()
    if line:
        set_line(s.line, line, line_w, line_alpha, dash)
    else:
        s.line.fill.background()
    return s


def glass(slide, x, y, w, h, radius=R_CARD, tint=GLASS, hair=GLASS_LINE):
    """A glass card: --foreground at a few percent over the canvas, with a 1px hairline."""
    return box(slide, x, y, w, h, FG, tint, radius, FG, hair)


def hline(slide, x0, x1, y, colour=FG, alpha=0.10, width_pt=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y), Inches(x1), Inches(y))
    set_line(c.line, colour, width_pt, alpha)
    return c


def shape_text(s, paras, size=16, color=FG, weight="regular", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
               margin=0.08, emph_color=KEY_INK, space_after=0):
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        txt, sz, col, wt = (item if isinstance(item, tuple) else (item, size, color, weight))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        _fill_runs(p, txt, sz, col, wt, emph_color)


# ---------------------------------------------------------------- slide chrome
def has_strip(spec):
    return spec.get("stage", "none") != "none"


def content_top(spec):
    return TOP_STRIP if has_strip(spec) else TOP_PLAIN


def process_strip(slide, steps, spec):
    """Breadcrumb above the title: current step filled in primary blue, past dim, future dimmer.
    steps is the section's own list: the six process steps in Part 1, the five chapters in the bonus section."""
    names = [p[0] for p in steps]
    stage, tail = spec["stage"], spec.get("stage_tail")
    cur = names.index(stage) if stage in names else -1
    y, h, gap = 0.30, 0.30, 0.14
    span = SW - 2 * MX
    tail_w = 1.6 if tail else 0
    cw = (span - gap * (len(names) - 1) - (gap + tail_w if tail else 0)) / len(names)
    for i, name in enumerate(names):
        x = MX + i * (cw + gap)
        label = f"{i + 1}  {name}"
        if i == cur:
            b = box(slide, x, y, cw, h, PRIMARY, radius=h / 2)
            shape_text(b, label, size=11, color=CANVAS, weight="semibold", margin=0.06)
        elif i < cur:
            b = box(slide, x, y, cw, h, FG, 0.06, h / 2, FG, 0.14)
            shape_text(b, label, size=11, color=MUTED, margin=0.06)
        else:
            b = box(slide, x, y, cw, h, None, None, h / 2, FG, 0.10)
            shape_text(b, label, size=11, color=DIM, margin=0.06)
    if tail:
        x = MX + len(names) * (cw + gap)
        b = box(slide, x, y, tail_w, h, None, None, h / 2, PRIMARY, None, 1.0)
        shape_text(b, f"→ {tail}", size=11, color=KEY, weight="semibold", margin=0.06)


def text_w(text, size, weight="regular"):
    """Width in inches of one line of text set in Inter (emphasis runs in semibold)."""
    total = 0.0
    for seg, b in _runs(text):
        f = _mfont("semibold" if b else weight)
        total += (f.getlength(seg) if f else len(seg) * 112) / 200 * size / 72
    return total


def chip_text(spec):
    """The small brand-tinted chip on the caption line: the comparison-table row a main slide illustrates,
    or the Part 1 slide a bonus slide builds on (rendered with its live number)."""
    if spec.get("finance"):
        return f"**In finance:** {spec['finance']}"
    if spec.get("seen"):
        return f"**Seen in Part 1:** slide {spec['seen_n']}"
    return None


def chrome(slide, spec, sections, total, steps):
    run = spec.get("run")
    strip = has_strip(spec)
    if strip:
        process_strip(slide, steps, spec)
    ty = 0.80 if strip else 0.38
    chip = chip_text(spec)
    narrow = run or chip or spec.get("optional")
    title_w = 8.55 if narrow else SW - 2 * MX
    tsize = 34 if est_lines(spec["title"], title_w, 34, "semibold") == 1 else 30
    add_text(slide, MX, ty, title_w, 1.25, spec["title"], size=tsize, weight="semibold", color=FG,
             space_after=0, tracking=-50)
    right = SW - MX
    cap_y = ty + 0.6
    if chip:
        cw = text_w(chip, 12, "regular") + 0.34
        c = box(slide, right - cw, cap_y - 0.02, cw, 0.32, BRAND, 0.20, 0.16, KEY, 0.55, 0.75)
        shape_text(c, chip, size=12, color=KEY_INK, emph_color=KEY, margin=0.12)
        cap_right = right - cw - 0.18
    else:
        cap_right = right
    if run:
        pill = box(slide, 9.45, ty + 0.02, 3.28, 0.5, PRIMARY, radius=0.25)
        shape_text(pill, f"▶ Notebook: run section {run}", size=15, color=CANVAS, weight="semibold")
        add_text(slide, cap_right - 4.6, cap_y, 4.6, 0.3, sections[run], size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    if spec.get("optional"):
        ow = text_w("optional", 12) + 0.36
        o = box(slide, right - ow, ty + 0.1, ow, 0.32, None, None, 0.16, FG, 0.30, 0.75)
        shape_text(o, "optional", size=12, color=MUTED, margin=0.1)
    if spec.get("reading"):
        ref = f"**{REF_LABEL}:** {spec['reading']}"
        ry = 6.52 if est_lines(ref, SW - 2 * MX, 11) == 1 else 6.44
        add_text(slide, MX, ry, SW - 2 * MX, 0.45, ref, size=11, color=MUTED, emph_color=FG, space_after=0)
    hline(slide, MX, SW - MX, 6.94, FG, 0.10)
    slide.shapes.add_picture(str(LOGO), Inches(MX), Inches(7.03), Inches(0.22), Inches(0.22))
    add_text(slide, MX + 0.32, 7.04, 6, 0.3, f"**{WORDMARK}** · {SOCIETY}", size=11, color=MUTED, emph_color=FG)
    where = "Bonus · " if spec["part"] == "bonus" else ""
    add_text(slide, SW - MX - 4, 7.04, 4, 0.3, f"{SITE}     {where}{spec['n']} / {total}", size=11, color=MUTED,
             align=PP_ALIGN.RIGHT)


def text_column(slide, x, y, w, spec, size=20, bottom=BOTTOM):
    """Definition cards first (terms are defined before they are used), then lines. Returns the y reached."""
    for term, text in spec.get("define", []):
        y = def_card(slide, x, y, w, term, text) + 0.22
    for line in spec.get("lines", []):
        h = text_h(est_lines(line, w, size), size)
        add_text(slide, x, y, w, h + 0.05, line, size=size)
        y += h + 0.22
    if y > bottom + 0.05:
        print(f"  warning: slide {spec['n']} text column may overflow (to {y:.2f} in)")
    return y


def def_card(slide, x, y, w, term, text, size=16):
    inner = w - 0.4
    body = f"**{term}:** {text}"
    h = text_h(est_lines(body, inner, size), size) + 0.3
    card = glass(slide, x, y, w, h)
    shape_text(card, body, size=size, color=FG, align=PP_ALIGN.LEFT, margin=0.2, emph_color=KEY)
    return y + h


# ---------------------------------------------------------------- visuals
def fit_picture(slide, path, x, y, w, h, align="center"):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    slide.shapes.add_picture(str(path), Inches(px), Inches(y), Inches(pw), Inches(ph))
    return ph


def v_image(slide, v, x, y, w, h):
    return fit_picture(slide, FIG / v["file"], x, y, w, h)


def v_result(slide, v, x, y, w, h):
    path = FIG / v["file"]
    cap_h = 0.38
    if path.exists():
        ph = fit_picture(slide, path, x, y, w, h - cap_h)
        add_text(slide, x, y + ph + 0.06, w, cap_h, v["caption"], size=12, color=MUTED,
                 align=PP_ALIGN.CENTER)
        return ph + cap_h
    s = box(slide, x, y, w, h, None, None, R_CARD, FG, 0.30, 1.0, dash=True)
    shape_text(s, [(f"{v['id']}: figure not generated yet", 20, MUTED, "semibold"),
                   (f"Run python3 deck/build_deck.py --figures to draw deck/figures/{v['file']}.", 13, MUTED, "regular"),
                   (v["caption"], 14, FG, "regular")], margin=0.3, space_after=6)
    return h


def v_stack(slide, v, x, y, w, h):
    """Items top to bottom. A concept image takes at most 40% of the height; result figures share the rest."""
    items = v["items"]
    gap = 0.2
    free = h - gap * (len(items) - 1)
    fixed = {}
    for i, it in enumerate(items):
        if it["kind"] == "image":
            with Image.open(FIG / it["file"]) as im:
                fixed[i] = min(free * 0.4, w * im.height / im.width)
    rest = (free - sum(fixed.values())) / max(1, len(items) - len(fixed))
    yy = y
    for i, it in enumerate(items):
        hh = fixed.get(i, rest)
        VISUALS[it["kind"]](slide, it, x, yy, w, hh)
        yy += hh + gap
    return h


def v_pipeline(slide, v, x, y, w, h):
    steps = v["steps"]
    n, gap = len(steps), 0.3
    bw = (w - gap * (n - 1)) / n
    inner = bw - 0.28
    ns = v.get("name_size", 20)
    while ns > 14 and any(text_w(max(name.split(" "), key=len), ns, "semibold") > inner * 0.88 for name, _ in steps):
        ns -= 1           # never break a word: shrink the step names until the longest word fits
    # measured at 88% of the box: LibreOffice Impress wraps these narrow captions earlier than Inter's metrics say
    need = max(0.2 + text_h(1, 13) + text_h(est_lines(name, inner * 0.88, ns, "semibold"), ns)
               + text_h(est_lines(cap, inner * 0.88, 13), 13) + 0.3 for name, cap in steps)
    bh = min(h, max(1.8, need))
    for i, (name, cap) in enumerate(steps):
        bx = x + i * (bw + gap)
        b = glass(slide, bx, y, bw, bh)
        shape_text(b, [(f"{i + 1}", 13, KEY, "semibold"), (name, ns, FG, "semibold"), (cap, 13, MUTED, "regular")],
                   anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margin=0.14, space_after=4)
        b.text_frame.margin_top = Inches(0.16)
        if i < n - 1:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(bx + bw + 0.07), Inches(y + bh / 2 - 0.09),
                                       Inches(gap - 0.14), Inches(0.18))
            a.shadow.inherit = False
            set_fill(a, RULE)
            a.line.fill.background()
    return bh


def v_cards(slide, v, x, y, w, h):
    cards, cols = v["cards"], v.get("cols", 1)
    rows = math.ceil(len(cards) / cols)
    gap = 0.22
    cw = (w - gap * (cols - 1)) / cols
    ch = min((h - gap * (rows - 1)) / rows, 1.6 if cols == 1 else 2.0)
    hs, bs = (20, 16) if cols == 1 else (18, 16)
    tile = 0.46 if cols == 1 else 0.4
    tx = 0.22 + tile + 0.18
    for k, (head, body) in enumerate(cards):
        cx = x + (k % cols) * (cw + gap)
        cy = y + (k // cols) * (ch + gap)
        glass(slide, cx, cy, cw, ch)
        t = box(slide, cx + 0.22, cy + 0.22, tile, tile, BRAND, 0.20, R_INPUT)
        shape_text(t, str(k + 1), size=15, color=KEY, weight="semibold", margin=0)
        tw = cw - tx - 0.2
        hh = text_h(est_lines(head, tw, hs, "semibold"), hs)
        add_text(slide, cx + tx, cy + 0.24, tw, hh + 0.05, head, size=hs, weight="semibold", color=FG, space_after=0)
        add_text(slide, cx + tx, cy + 0.32 + hh, tw, ch - hh - 0.42, body, size=bs, color=MUTED, space_after=0)
    return rows * ch + (rows - 1) * gap


def data_stats():
    if not DATA.exists():
        return {"n_days": "about 5,000", "n_rows": "about 200,000"}
    days, rows = set(), 0
    with gzip.open(DATA, "rt", newline="") as fh:
        for r in csv.DictReader(fh):
            rows += 1
            days.add(r["date"])
    return {"n_days": f"{len(days):,}", "n_rows": f"{rows:,}"}


def v_table(slide, v, x, y, w, h):
    """The data card: one glass card, a row per fact, hairlines between rows, the warning row in red."""
    rows = v["rows"]
    stats = data_stats()
    k0, pad, size = 1.75, 0.22, 16
    vw = w - k0 - pad
    heights = [text_h(est_lines(val.format(**stats), vw, size), size) + 0.26 for _, val in rows]
    total = sum(heights)
    glass(slide, x, y, w, total)
    yy = y
    for i, ((key, val), rh) in enumerate(zip(rows, heights)):
        warn = i == v.get("warn_row")
        if warn:
            box(slide, x + 0.06, yy + 0.04, w - 0.12, rh - 0.08, WARN, 0.10, R_INPUT)
        col = WARN if warn else FG
        add_text(slide, x + pad, yy, k0 - pad - 0.1, rh, key, size=15, weight="semibold",
                 color=WARN if warn else KEY, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        add_text(slide, x + k0, yy, vw, rh, val.format(**stats), size=size, color=col,
                 weight="semibold" if warn else "regular", anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        if i < len(rows) - 1 and not warn and i + 1 != v.get("warn_row"):
            hline(slide, x + pad, x + w - pad, yy + rh, FG, 0.10)
        yy += rh
    return total


def v_stat(slide, v, x, y, w, h):
    bh = min(h, 3.4)
    b = glass(slide, x + 0.3, y + 0.2, w - 0.6, bh)
    shape_text(b, [(v["big"], v.get("big_size", 130), KEY_INK, "semibold"), (v["label"], v.get("label_size", 22), FG,
                                                                             "regular")], margin=0.3)
    return bh


def _no_fill_sppr():
    sp = etree.Element(qn("c:spPr"))
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    return sp


def _style_chart(chart, size=12):
    """Dark chart: transparent chart and plot areas over the canvas, Inter, muted text."""
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(MUTED)
    cs = chart._chartSpace
    for old in cs.findall(qn("c:spPr")):
        cs.remove(old)
    txpr = cs.find(qn("c:txPr"))
    (txpr.addprevious if txpr is not None else cs.append)(_no_fill_sppr())
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    for old in pa.findall(qn("c:spPr")):
        pa.remove(old)
    ext = pa.find(qn("c:extLst"))
    (ext.addprevious if ext is not None else pa.append)(_no_fill_sppr())


def _axis_title(axis, text):
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = text
    style_font(tf.paragraphs[0].runs[0].font, 13, MUTED)


def _gridlines(axis, on=True):
    axis.has_major_gridlines = on
    if on:
        set_line(axis.major_gridlines.format.line, GRID, 0.75)
    set_line(axis.format.line, RULE, 0.75)


def _monthly_prices(symbols):
    """First trading day of each month, adj_close, rescaled to 100 at the first point."""
    out = {s: [] for s in symbols}
    seen = set()
    with gzip.open(DATA, "rt", newline="") as fh:
        for r in csv.DictReader(fh):
            s = r["symbol"]
            if s not in out:
                continue
            ym = (s, r["date"][:7])
            if ym in seen:
                continue
            seen.add(ym)
            yr, mo = int(r["date"][:4]), int(r["date"][5:7])
            out[s].append((yr + (mo - 1) / 12, float(r["adj_close"])))
    for s in symbols:
        pts = sorted(out[s])
        base = pts[0][1]
        out[s] = [(t, 100 * p / base) for t, p in pts]
    return out


def v_chart_prices(slide, v, x, y, w, h):
    cap_h = 0.4
    if not DATA.exists():
        return v_result(slide, {"id": "Price chart", "file": "../../data/prices.csv.gz", "caption": v["caption"]}, x, y, w, h)
    series = _monthly_prices(v["symbols"])
    cd = XyChartData()
    for s, pts in series.items():
        ser = cd.add_series(s)
        for t, p in pts:
            ser.add_data_point(t, round(p, 2))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y), Inches(w),
                                Inches(h - cap_h), cd)
    ch = gf.chart
    _style_chart(ch)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    style_font(ch.legend.font, 13, FG)
    looks = [(KEY, False), (FG, False), (AMBER, False), (MUTED, True)]   # blue, off-white, amber, grey dashed
    for i, ser in enumerate(ch.series):
        colour, dash = looks[i % len(looks)]
        set_line(ser.format.line, colour, 2.25, dash=dash)
        ser.smooth = False
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = 2000, 2020, 5
    xa.tick_labels.number_format, xa.tick_labels.number_format_is_linked = "0", False
    _gridlines(xa, False)
    _gridlines(ya, True)
    ya.minimum_scale = 0
    _axis_title(ya, "price, start = 100")
    add_text(slide, x, y + h - cap_h + 0.05, w, cap_h, v["caption"], size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return h


def v_chart_scatter(slide, v, x, y, w, h):
    cap_h = 0.4
    rnd = random.Random(11)
    r = v["r"]
    cd = XyChartData()
    ser = cd.add_series("synthetic points")
    for _ in range(v["n"]):
        a = rnd.gauss(0, 1)
        b = r * a + math.sqrt(1 - r * r) * rnd.gauss(0, 1)
        ser.add_data_point(round(a, 3), round(b, 3))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x), Inches(y), Inches(w), Inches(h - cap_h), cd)
    ch = gf.chart
    _style_chart(ch)
    ch.has_legend = False
    s = ch.series[0]
    s.format.line.fill.background()
    s.marker.style = XL_MARKER_STYLE.CIRCLE
    s.marker.size = 5
    s.marker.format.fill.solid()
    s.marker.format.fill.fore_color.rgb = rgb(KEY)
    s.marker.format.line.fill.background()
    for ax, t in ((ch.category_axis, "forecast"), (ch.value_axis, "what actually happened")):
        ax.minimum_scale, ax.maximum_scale = -3.5, 3.5
        ax.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        _gridlines(ax, False)
        _axis_title(ax, t)
    add_text(slide, x, y + h - cap_h + 0.05, w, cap_h, v["caption"], size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return h


def expected_max_sharpe(n, years):
    """Bailey et al. (2014) approximation of E[max] of n zero-skill annualised Sharpe ratios over `years`."""
    if n <= 1:
        return 0.0
    g = 0.5772156649
    z = NormalDist().inv_cdf
    return ((1 - g) * z(1 - 1 / n) + g * z(1 - 1 / (n * math.e))) / math.sqrt(years)


def v_chart_trials(slide, v, x, y, w, h):
    cap_h = 0.45
    T, N = v["years"], v["max_n"]
    cd = XyChartData()
    s1 = cd.add_series("expected best Sharpe, zero skill")
    for n in range(1, N + 1):
        s1.add_data_point(n, round(expected_max_sharpe(n, T), 3))
    s2 = cd.add_series("Sharpe = 1")
    s2.add_data_point(1, 1.0)
    s2.add_data_point(N, 1.0)
    s3 = cd.add_series("7 tries")
    s3.add_data_point(7, round(expected_max_sharpe(7, T), 3))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y), Inches(w),
                                Inches(h - cap_h), cd)
    ch = gf.chart
    _style_chart(ch)
    ch.has_legend = False
    a, b, c = ch.series
    set_line(a.format.line, KEY, 3)
    set_line(b.format.line, MUTED, 1.5, dash=True)
    c.marker.style, c.marker.size = XL_MARKER_STYLE.CIRCLE, 12
    c.marker.format.fill.solid()
    c.marker.format.fill.fore_color.rgb = rgb(WARN)
    c.marker.format.line.fill.background()
    # The annotation is a text box, not a data label: data labels wrap into three lines over the curve.
    # The marker sits at x=7 of 0..N and y=1.0 of 0..2 inside the plot area; the plot area is roughly
    # the chart frame inset by the axis titles (left ~0.75in, bottom ~0.75in, top ~0.15in).
    px = x + 0.75 + (w - 0.95) * (7 / N) + 0.45
    py = y + 0.15 + ((h - cap_h) - 0.9) * (1 - 1.0 / 2.0) - 0.45
    add_text(slide, px, py, 3.2, 0.4, f"**7 tries: Sharpe {expected_max_sharpe(7, T):.1f}, by luck alone**",
             size=14, color=WARN, weight="semibold", emph_color=WARN, space_after=0)
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, N, 10
    ya.minimum_scale, ya.maximum_scale, ya.major_unit = 0, 2.0, 0.5
    ya.tick_labels.number_format, ya.tick_labels.number_format_is_linked = "0.0", False
    _gridlines(xa, False)
    _gridlines(ya, True)
    _axis_title(xa, "number of strategies tried")
    _axis_title(ya, "best Sharpe ratio in the backtest")
    add_text(slide, x, y + h - cap_h + 0.05, w, cap_h, v["caption"], size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return h


def v_portfolio(slide, v, x, y, w, h):
    n, nl, ns = v["n"], v["n_long"], v["n_short"]
    add_text(slide, x, y, w, 0.35, "Stocks ranked by forecast (illustrative numbers)", size=14, color=MUTED,
             align=PP_ALIGN.LEFT)
    y0 = y + 0.45
    rh = (h - 0.5) / n
    label_w, bar_w = 1.05, 2.35
    centre = x + label_w + bar_w / 2
    fc = [0.8 - i * (1.6 / (n - 1)) for i in range(n)]
    for i in range(n):
        role = "long" if i < nl else ("short" if i >= n - ns else "none")
        col = {"long": BRAND, "short": MUTED, "none": NONE_BAR}[role]
        ry = y0 + i * rh
        add_text(slide, x, ry + 0.02, label_w - 0.1, rh, f"Stock {chr(65 + i)}", size=12, color=FG,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        length = abs(fc[i]) / 0.8 * (bar_w / 2 - 0.05)
        bx = centre if fc[i] >= 0 else centre - length
        if length > 0.01:
            box(slide, bx, ry + rh * 0.18, length, rh * 0.64, col, radius=0.03)
        add_text(slide, centre + bar_w / 2 + 0.08, ry, 0.7, rh, f"{fc[i]:+.1f}%", size=12, color=MUTED,
                 anchor=MSO_ANCHOR.MIDDLE, space_after=0, align=PP_ALIGN.LEFT)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(centre), Inches(y0), Inches(centre),
                                    Inches(y0 + n * rh))
    set_line(ln.line, RULE, 1.0)
    gx = x + label_w + bar_w + 0.85
    gw = w - (gx - x)
    groups = [(0, nl, KEY, "Buy (long)", f"weight +1/{nl} each"),
              (nl, n - ns, MUTED, "No position", "weight 0"),
              (n - ns, n, FG, "Sell short", f"weight -1/{ns} each")]
    for a, b, col, head, sub in groups:
        add_text(slide, gx, y0 + a * rh, gw, (b - a) * rh, [f"**{head}**", sub], size=16, color=col,
                 emph_color=col, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    return h


def v_qr(slide, v, x, y, w, h):
    codes = make_qr.make_all()
    n, gap = len(codes), 0.4
    q = min((w - gap * (n - 1)) / n, 1.8)
    total = n * q + (n - 1) * gap
    x0 = x + (w - total) / 2
    for i, c in enumerate(codes):
        qx = x0 + i * (q + gap)
        box(slide, qx - 0.07, y - 0.07, q + 0.14, q + 0.14, WHITE, radius=R_INPUT)   # light quiet zone
        slide.shapes.add_picture(str(c["path"]), Inches(qx), Inches(y), Inches(q), Inches(q))
        label = c["label"] + (f"\n{SITE}" if c["key"] == "website" and not c["placeholder"] else "")
        add_text(slide, qx - 0.15, y + q + 0.14, q + 0.3, 0.7, label.split("\n"), size=14, color=FG,
                 align=PP_ALIGN.CENTER, space_after=0)
        if c["placeholder"]:
            box(slide, qx - 0.1, y - 0.1, q + 0.2, q + 0.2, None, None, R_INPUT, WARN, None, 3)
            tag = box(slide, qx - 0.1, y + q / 2 - 0.36, q + 0.2, 0.72, WARN, radius=R_INPUT)
            shape_text(tag, [("PLACEHOLDER", 14, CANVAS, "semibold"), ("replace before use", 12, CANVAS, "regular")],
                       margin=0.04)
    add_text(slide, x, y + q + 0.9, w, 0.4, "Scan to follow QuantSoc", size=16, weight="semibold", color=KEY,
             align=PP_ALIGN.CENTER)
    return h


def v_grid(slide, v, x, y, w, h):
    """A general table on one glass card: an optional header row, then rows; the first column is the row key
    (brand-400 semibold). `widths` are column fractions; `highlight` tints one column (e.g. the finance column).
    Row heights follow the tallest cell, measured with Inter, so long cells wrap instead of overflowing."""
    rows, header = v["rows"], v.get("header")
    ncol = len(rows[0])
    fr = v.get("widths") or [1 / ncol] * ncol
    cws = [w * f / sum(fr) for f in fr]
    size, hsize = v.get("size", 15), v.get("header_size", v.get("size", 15) - 1)
    padx, pady = v.get("pad_x", 0.16), v.get("pad_y", 0.1)
    hi = v.get("highlight")

    def row_h(cells, sz, weights):
        return max(text_h(est_lines(c, cw - 2 * padx, sz, wt), sz) for c, cw, wt in zip(cells, cws, weights)) + 2 * pady

    body_w = ["semibold"] + ["regular"] * (ncol - 1)
    hh = row_h(header, hsize, ["semibold"] * ncol) if header else 0
    rhs = [row_h(r, size, body_w) for r in rows]
    total = hh + sum(rhs)
    if total > h + 0.05:
        print(f"  warning: grid needs {total:.2f} in, has {h:.2f} in")
    glass(slide, x, y, w, total)
    if hi is not None:
        hx = x + sum(cws[:hi])
        box(slide, hx + 0.04, y + 0.04, cws[hi] - 0.08, total - 0.08, BRAND, 0.10, R_INPUT)
    yy = y
    if header:
        cx = x
        for k, (c, cw) in enumerate(zip(header, cws)):
            add_text(slide, cx + padx, yy, cw - 2 * padx, hh, c, size=hsize, weight="semibold",
                     color=KEY if k == hi or k == 0 else FG, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
            cx += cw
        yy += hh
        hline(slide, x + padx, x + w - padx, yy, FG, 0.22)
    for i, (r, rh) in enumerate(zip(rows, rhs)):
        cx = x
        for k, (c, cw) in enumerate(zip(r, cws)):
            add_text(slide, cx + padx, yy, cw - 2 * padx, rh, c, size=size, weight="semibold" if k == 0 else "regular",
                     color=KEY if k == 0 else FG, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
            cx += cw
        yy += rh
        if i < len(rows) - 1:
            hline(slide, x + padx, x + w - padx, yy, FG, 0.10)
    return total


def v_pair(slide, v, x, y, w, h):
    """Two figures side by side (result figures keep their placeholder when the PNG is missing)."""
    gap = 0.3
    cw = (w - gap) / 2
    used = 0
    for k, it in enumerate(v["items"]):
        used = max(used, VISUALS[it["kind"]](slide, it, x + k * (cw + gap), y, cw, h))
    return used


def v_tiles(slide, v, x, y, w, h):
    """A grid of glass cards, each {head, body, foot, mark}: head in semibold, body in off-white, foot muted.
    A marked card gets a primary hairline and a filled badge (v['mark_label']) straddling its top edge.
    All cards share one height, the tallest card's measured need."""
    cards, cols = v["cards"], v.get("cols", 4)
    rows = math.ceil(len(cards) / cols)
    gap = v.get("gap", 0.22)
    top_pad = 0.2 if any(c.get("mark") for c in cards) else 0
    cw = (w - gap * (cols - 1)) / cols
    hs, bs, fs = v.get("head_size", 16), v.get("body_size", 13), v.get("foot_size", 12)
    inner = cw - 0.36

    pad = v.get("pad", 0.16)

    def need(c):
        n = pad + text_h(est_lines(c["head"], inner, hs, "semibold"), hs) + 0.06
        n += text_h(est_lines(c.get("body", ""), inner, bs), bs) if c.get("body") else 0
        n += (0.06 + text_h(est_lines(f"**Ours:** {c['ours']}", inner, bs), bs)) if c.get("ours") else 0
        n += (0.08 + text_h(est_lines(c["foot"], inner, fs), fs)) if c.get("foot") else 0
        return n + pad
    row_h = [max(need(c) for c in cards[r * cols:(r + 1) * cols]) for r in range(rows)]
    if v.get("even", True):
        row_h = [max(row_h)] * rows
    total = top_pad + sum(row_h) + (rows - 1) * gap
    if total > h + 0.05:
        print(f"  warning: tiles need {total:.2f} in, have {h:.2f} in")
    for k, c in enumerate(cards):
        r = k // cols
        ch = row_h[r]
        cx = x + (k % cols) * (cw + gap)
        cy = y + top_pad + sum(row_h[:r]) + r * gap
        if c.get("mark"):
            box(slide, cx, cy, cw, ch, FG, GLASS, R_CARD, PRIMARY, None, 1.5)
            label = v.get("mark_label", "used here")
            bw = text_w(label, 12, "semibold") + 0.3
            b = box(slide, cx + cw - bw - 0.16, cy - 0.15, bw, 0.3, PRIMARY, radius=0.15)
            shape_text(b, label, size=12, color=CANVAS, weight="semibold", margin=0.06)
        else:
            glass(slide, cx, cy, cw, ch)
        yy = cy + pad
        hh = text_h(est_lines(c["head"], inner, hs, "semibold"), hs)
        add_text(slide, cx + 0.18, yy, inner, hh + 0.05, c["head"], size=hs, weight="semibold", color=KEY,
                 space_after=0)
        yy += hh + 0.06
        if c.get("body"):
            bh = text_h(est_lines(c["body"], inner, bs), bs)
            add_text(slide, cx + 0.18, yy, inner, bh + 0.05, c["body"], size=bs, color=FG, space_after=0)
            yy += bh
        if c.get("ours"):
            oh = text_h(est_lines(f"**Ours:** {c['ours']}", inner, bs), bs)
            add_text(slide, cx + 0.18, yy + 0.06, inner, oh + 0.05, f"**Ours:** {c['ours']}", size=bs, color=KEY_INK,
                     emph_color=KEY_INK, space_after=0)
            yy += oh + 0.06
        if c.get("foot"):
            fh = text_h(est_lines(c["foot"], inner, fs), fs)
            add_text(slide, cx + 0.18, cy + ch - pad - fh, inner, fh + 0.05, c["foot"], size=fs, color=MUTED,
                     space_after=0)
    return total


def v_sets(slide, v, x, y, w, h):
    """Train / validation / test drawn natively: one time bar cut into blocks by share, then a card per block
    saying its job. Train in primary, validation as a brand tile, the test block as a raised neutral."""
    sets = v["sets"]       # [name, job, text, share]
    looks = [(PRIMARY, None, CANVAS), (BRAND, 0.20, KEY), (FG, 0.10, FG)]
    bar_y, bar_h, gap = y + 0.05, 0.62, 0.08
    tot = sum(s[3] for s in sets)
    span = w - gap * (len(sets) - 1)
    bx = x
    for k, (name, job, text, share) in enumerate(sets):
        fill, a, ink = looks[k % len(looks)]
        bw = span * share / tot
        b = box(slide, bx, bar_y, bw, bar_h, fill, a, R_INPUT)
        shape_text(b, f"**{name}**", size=17, color=ink, emph_color=ink, margin=0.1)
        bx += bw + gap
    ay = bar_y + bar_h + 0.2
    a = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(ay), Inches(x + w - 0.9), Inches(ay))
    set_line(a.line, RULE, 1.25)
    a.line._get_or_add_ln().append(etree.fromstring(
        '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle"/>'))
    add_text(slide, x + w - 0.8, ay - 0.15, 0.8, 0.3, "time", size=13, color=MUTED, space_after=0)
    cy = ay + 0.3
    cg = 0.22
    cw = (w - cg * (len(sets) - 1)) / len(sets)
    inner = cw - 0.4
    ch = max(0.22 + text_h(1, 18) + 0.1 + text_h(est_lines(t, inner, 15), 15) + 0.22 for _, _, t, _ in sets)
    for k, (name, job, text, share) in enumerate(sets):
        cx = x + k * (cw + cg)
        glass(slide, cx, cy, cw, ch)
        add_text(slide, cx + 0.2, cy + 0.22, inner, text_h(1, 18) + 0.05, f"{name}: **{job}**", size=18,
                 weight="semibold", color=FG, emph_color=KEY, space_after=0)
        add_text(slide, cx + 0.2, cy + 0.32 + text_h(1, 18), inner, ch - 0.5, text, size=15, color=MUTED,
                 space_after=0)
    return cy + ch - y


VISUALS = {"image": v_image, "result": v_result, "stack": v_stack, "pipeline": v_pipeline, "cards": v_cards,
           "table": v_table, "stat": v_stat, "chart_prices": v_chart_prices, "chart_scatter": v_chart_scatter,
           "chart_trials": v_chart_trials, "portfolio": v_portfolio, "qr": v_qr, "grid": v_grid, "pair": v_pair,
           "tiles": v_tiles, "sets": v_sets}


# ---------------------------------------------------------------- layouts
def paint_canvas(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(CANVAS)


def layout_title(slide, spec):
    slide.shapes.add_picture(str(LOGO), Inches(0.8), Inches(0.62), Inches(0.82), Inches(0.82))
    add_text(slide, 1.85, 0.6, 6, 0.5, WORDMARK, size=30, weight="semibold", color=FG, space_after=0, tracking=-40)
    add_text(slide, 1.85, 1.12, 6, 0.4, SOCIETY, size=17, color=MUTED, space_after=0)
    add_text(slide, 0.8, 2.3, 7.4, 2.2, spec["title"], size=48, weight="bold", color=FG, space_after=0,
             tracking=-100)
    add_text(slide, 0.8, 4.55, 7.4, 0.5, spec["subtitle"], size=22, color=MUTED)
    note = glass(slide, 0.8, 5.3, 7.6, 1.1)
    shape_text(note, spec["lines"][0], size=17, color=FG, weight="semibold", align=PP_ALIGN.LEFT, margin=0.25)
    add_text(slide, 0.8, 6.8, 6, 0.4, SITE, size=18, weight="semibold", color=KEY)
    # motif: ten rising decile bars, the picture this workshop builds towards, brand blue gaining opacity
    bx, base, bw, gap = 8.7, 6.2, 0.3, 0.12
    for i in range(10):
        hh = 0.9 + i * 0.36 + (0.25 if i in (3, 6) else 0) - (0.2 if i in (4, 8) else 0)
        box(slide, bx + i * (bw + gap), base - hh, bw, hh, BRAND, 0.28 + 0.08 * i, R_INPUT)
    hline(slide, bx - 0.1, bx + 10 * (bw + gap) - gap + 0.1, base + 0.06, FG, 0.18)


def layout_side(slide, spec):
    tw, vx = 5.35, 6.3
    top = content_top(spec)
    text_column(slide, MX, top, tw, spec, size=spec.get("size", 20))
    v = spec.get("visual")
    if v:
        VISUALS[v["kind"]](slide, v, vx, top, SW - MX - vx, BOTTOM - top)


def layout_wide(slide, spec):
    v = spec["visual"]
    w = SW - 2 * MX
    top = content_top(spec)
    if v["kind"] == "image":
        vh = fit_picture(slide, FIG / v["file"], MX, top - 0.1, w, 2.45 if has_strip(spec) else 2.55)
    else:
        vh = VISUALS[v["kind"]](slide, v, MX, top, w, 3.0)
    y = top + vh + (0.12 if v["kind"] == "image" else 0.28)
    defs, lines = spec.get("define", []), spec.get("lines", [])
    if defs and lines:
        lw = (w - 0.4) / 2
        yy = y
        for term, text in defs:
            yy = def_card(slide, MX, yy, lw, term, text) + 0.15
        text_column(slide, MX + lw + 0.4, y, lw, {"n": spec["n"], "lines": lines})
    else:
        text_column(slide, MX, y, w, spec, size=spec.get("size", 22))
    chips = spec.get("chips", [])
    for k, chip in enumerate(chips):
        cwid = (w - 0.2 * (len(chips) - 1)) / len(chips)
        c = glass(slide, MX + k * (cwid + 0.2), 5.52, cwid, 0.85)
        shape_text(c, [(f"Task {k + 1}", 13, KEY, "semibold"), (chip, 15, FG, "regular")], margin=0.1)


def layout_full(slide, spec):
    """The visual across the full width (a grid, tiles, a pair of figures), then any lines beneath it."""
    w = SW - 2 * MX
    top = content_top(spec)
    size = spec.get("size", 18)
    lines = spec.get("lines", [])
    need = sum(text_h(est_lines(t, w, size), size) + 0.12 for t in lines)
    room = BOTTOM - top - (need + 0.18 if lines else 0)
    vh = VISUALS[spec["visual"]["kind"]](slide, spec["visual"], MX, top, w, room)
    y = top + vh + (0.2 if lines else 0)
    for t in lines:
        th = text_h(est_lines(t, w, size), size)
        add_text(slide, MX, y, w, th + 0.05, t, size=size)
        y += th + 0.12
    if y > BOTTOM + 0.08:
        print(f"  warning: slide {spec['n']} ({spec['id']}) full layout may overflow (to {y:.2f} in)")


def layout_figure(slide, spec):
    """One dense result figure as large as the slide allows, with at most two short lines. The figure carries its
    own title, so no caption is drawn. A 16:10 figure is height-capped on this canvas, so the figure fills the full
    content height flush right with the lines in a narrow column on the left; a figure wide enough to gain from
    the full width instead goes full width with the lines underneath. Whichever placement draws it larger wins."""
    v = spec["visual"]
    path = FIG / v["file"]
    lines, size = spec.get("lines", []), spec.get("size", 18)
    if len(lines) > 2:
        print(f"  warning: slide {spec['n']} ({spec['id']}) figure layout takes at most two lines")
    top = content_top(spec) - 0.1
    bottom = BOTTOM if spec.get("reading") else 6.8
    w = SW - 2 * MX
    if not path.exists():
        layout_side(slide, spec)
        return
    with Image.open(path) as im:
        aspect = im.width / im.height
    tw, gap = 4.0, 0.3
    beside = min((w - tw - gap) / aspect, bottom - top)
    below_lines = sum(text_h(est_lines(t, w, size), size) + 0.12 for t in lines)
    below = min(w / aspect, bottom - top - below_lines - 0.12)
    if beside >= below:
        ph = beside
        pw = ph * aspect
        slide.shapes.add_picture(str(path), Inches(SW - MX - pw), Inches(top), Inches(pw), Inches(ph))
        tcol = SW - MX - pw - gap - MX
        y = top + 0.1
        for t in lines:
            th = text_h(est_lines(t, tcol, size), size)
            add_text(slide, MX, y, tcol, th + 0.05, t, size=size)
            y += th + 0.25
        if y > bottom + 0.05:
            print(f"  warning: slide {spec['n']} ({spec['id']}) figure text may overflow (to {y:.2f} in)")
    else:
        pw = below * aspect
        slide.shapes.add_picture(str(path), Inches(MX + (w - pw) / 2), Inches(top), Inches(pw), Inches(below))
        y = top + below + 0.12
        for t in lines:
            th = text_h(est_lines(t, w, size), size)
            add_text(slide, MX, y, w, th + 0.05, t, size=size)
            y += th + 0.12


LAYOUTS = {"title": layout_title, "side": layout_side, "wide": layout_wide, "full": layout_full,
           "figure": layout_figure}


# ---------------------------------------------------------------- figures
def ensure_result_figures(force=False):
    """Draw F1..F4 on the dark canvas when one is missing or force is set; otherwise keep the cached PNGs."""
    missing = [f for f in RESULT_FIGS if not (FIG / f).exists()]
    if force or missing:
        try:
            sys.path.insert(0, str(ROOT))
            from matplotlib import font_manager
            for ttf in FONT_DIR.glob("*.ttf"):
                font_manager.fontManager.addfont(str(ttf))
            from quantsoc_ml.data import daily_returns, load_prices, to_wide
            from quantsoc_ml.notebook_support import make_deck_figures
            long = load_prices(DATA)
            prices, volume = to_wide(long, "adj_close"), to_wide(long, "volume")
            print("drawing F1..F4 for the dark canvas (about 30 seconds)")
            make_deck_figures(FIG, prices, volume, daily_returns(prices), dark=True)
        except Exception as exc:  # the deck must still build: placeholders are drawn for what is missing
            print(f"  warning: could not draw the result figures ({type(exc).__name__}: {exc})")
    for f in RESULT_FIGS:
        p = FIG / f
        if p.exists():
            with Image.open(p) as im:
                if sum(im.convert("RGB").getpixel((2, 2))) > 120:
                    print(f"  warning: {f} has a light background; rebuild with --figures for the dark version")


# ---------------------------------------------------------------- notes, checks, build
def notes_text(spec):
    a, b = spec["minutes"]
    label = "Bonus minutes" if spec["part"] == "bonus" else "Minutes"
    parts = [f"{label} {a} to {b} ({b - a} min)."]
    if spec.get("optional"):
        parts.append("Optional slide: skip it when time is short; nothing later depends on it.")
    parts.append(spec["notes"].strip())
    if spec.get("cue"):
        parts.append("Cue: " + spec["cue"])
    if spec.get("reading_full"):
        parts.append("Full citation: " + " ".join(spec["reading_full"]))
    return "\n\n".join(parts)


DASHES = re.compile("[—–]")
SLIDE_REF = re.compile(r"(?<!\{)\{slide:([a-z0-9_]+)\}(?!\})")


def load_spec():
    """Both yaml files as one slide list: Part 1 first, then the bonus section. Each slide gets its part, its
    live number n, and every "{slide:<id>}" in its text replaced by that slide's number. Returns
    (slides, sections, process, chapters, errors)."""
    main = yaml.safe_load((DECK / "slides.yaml").read_text(encoding="utf-8"))
    bonus_path = DECK / "slides_bonus.yaml"
    bonus = yaml.safe_load(bonus_path.read_text(encoding="utf-8")) if bonus_path.exists() else {}
    errs = []
    slides = []
    for part, src in (("main", main.get("slides", [])), ("bonus", bonus.get("slides", []))):
        for s in src:
            s["part"] = part
            slides.append(s)
    ids = {}
    for k, s in enumerate(slides, 1):
        s["n"] = k
        if "id" not in s:
            errs.append(f"slide {k}: no id")
        elif s["id"] in ids:
            errs.append(f"slide {k}: duplicate id {s['id']!r}")
        else:
            ids[s["id"]] = k

    def sub(obj, where):
        if isinstance(obj, str):
            def rep(m):
                if m.group(1) not in ids:
                    errs.append(f"{where}: unknown slide reference {{slide:{m.group(1)}}}")
                    return m.group(0)
                return str(ids[m.group(1)])
            return SLIDE_REF.sub(rep, obj)
        if isinstance(obj, list):
            return [sub(o, where) for o in obj]
        if isinstance(obj, dict):
            return {k: sub(v, where) for k, v in obj.items()}
        return obj
    for k, s in enumerate(slides):
        slides[k] = s = sub(s, f"slide {s['n']} ({s.get('id')})")
        if s.get("seen"):
            s["seen_n"] = ids.get(s["seen"])
    process, chapters = main["process"], bonus.get("chapters", [])
    for s in slides:
        v = s.get("visual") or {}
        if v.get("kind") == "pipeline" and "steps" not in v:
            v["steps"] = process if s["part"] == "main" else chapters
    sections = {int(k): v for k, v in main["sections"].items()}
    return slides, sections, process, chapters, errs


def finance_rows(slides):
    for s in slides:
        if s.get("id") == FINANCE_TABLE:
            return [r[0] for r in s["visual"]["rows"]]
    return []


def self_check(slides, process, chapters):
    errs = []
    main = [s for s in slides if s["part"] == "main"]
    bonus = [s for s in slides if s["part"] == "bonus"]
    if len(main) != MAIN_COUNT:
        errs.append(f"expected {MAIN_COUNT} Part 1 slides, found {len(main)}")
    if len(bonus) != BONUS_COUNT:
        errs.append(f"expected {BONUS_COUNT} bonus slides, found {len(bonus)}")
    if [p[0] for p in process] != PROCESS_STEPS:
        errs.append(f"process steps are {[p[0] for p in process]}")
    if [c[0] for c in chapters] != BONUS_CHAPTERS:
        errs.append(f"bonus chapters are {[c[0] for c in chapters]}")
    rows = finance_rows(slides)
    if not rows:
        errs.append(f"no comparison table on slide id {FINANCE_TABLE!r}")
    main_ids = {s["id"] for s in main}
    for part, group, expect in (("Part 1", main, MAIN_MINUTES), ("bonus", bonus, BONUS_MINUTES)):
        t = 0
        for s in group:
            a, b = s["minutes"]
            if a != t or b <= a:
                errs.append(f"slide {s['n']} ({s['id']}): minutes {a}-{b} not contiguous from {t}")
            t = b
        if t != expect:
            errs.append(f"{part} minute budget ends at {t}, expected {expect}")
    for s in slides:
        tag = f"slide {s['n']} ({s.get('id')})"
        if s.get("run") != EXPECTED_RUNS.get(s.get("id")):
            errs.append(f"{tag}: run marker {s.get('run')} expected {EXPECTED_RUNS.get(s.get('id'))}")
        if s.get("run") and "▶ Notebook: run section" not in s.get("cue", ""):
            errs.append(f"{tag}: run slide without a notebook cue in its notes")
        if "stage" not in s:
            errs.append(f"{tag}: no stage key")
        elif s["stage"] != EXPECTED_STAGES.get(s.get("id")):
            errs.append(f"{tag}: stage {s['stage']!r} expected {EXPECTED_STAGES.get(s.get('id'))!r}")
        if s.get("finance") and s["finance"] not in rows:
            errs.append(f"{tag}: 'In finance' chip names {s['finance']!r}, not a row of the table {rows}")
        if s.get("finance") and s["part"] != "main":
            errs.append(f"{tag}: 'In finance' chips belong to Part 1")
        if s.get("seen") and (s["part"] != "bonus" or s["seen"] not in main_ids):
            errs.append(f"{tag}: 'Seen in Part 1' must be on a bonus slide and name a Part 1 id, got {s['seen']!r}")
        if s.get("optional") and s["part"] != "bonus":
            errs.append(f"{tag}: only bonus slides are optional")
        n_text = len(s.get("lines", [])) + len(s.get("define", []))
        if n_text > 5:
            errs.append(f"{tag}: {n_text} text items (max 5)")
        blob = yaml.safe_dump(s, allow_unicode=True)
        if DASHES.search(blob):
            errs.append(f"{tag}: contains an em or en dash")
        if "{{bonus:" in yaml.safe_dump({k: v for k, v in s.items() if k != "notes"}, allow_unicode=True):
            errs.append(f"{tag}: a {{{{bonus:KEY}}}} placeholder outside the notes (slide text never quotes B1..B9)")
    if DASHES.search(yaml.safe_dump([process, chapters], allow_unicode=True)):
        errs.append("process steps or chapters contain an em or en dash")
    for what, fg, bg, floor in CONTRAST_PAIRS:
        if contrast(fg, bg) < floor:
            errs.append(f"contrast: {what} #{fg} on #{bg} is {contrast(fg, bg):.2f}:1, floor {floor}:1")
    return errs


def check_built(prs, slides):
    """After saving: every slide with a reference carries a 'Reference:' footer, no 'Reading:' label is left,
    no dash and no unrendered {slide:..} or {{bonus:..}} template reached a slide."""
    errs = []
    for s, sl in zip(slides, prs.slides):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        n_ref = sum(t.startswith(f"{REF_LABEL}:") for t in texts)
        if n_ref != (1 if s.get("reading") else 0):
            errs.append(f"slide {s['n']}: {n_ref} '{REF_LABEL}:' footers")
        if any("Reading:" in t for t in texts):
            errs.append(f"slide {s['n']}: a 'Reading:' label is still on the slide")
        if any(DASHES.search(t) for t in texts):
            errs.append(f"slide {s['n']}: an em or en dash reached the slide")
        if any("{slide:" in t or "{{" in t for t in texts):
            errs.append(f"slide {s['n']}: an unrendered template reached the slide")
        if SLIDE_REF.search(sl.notes_slide.notes_text_frame.text):
            errs.append(f"slide {s['n']}: an unrendered slide reference in the notes")
    return errs


def write_notes_md(slides, sections):
    out = ["# Speaker notes: Machine Learning for Quantitative Finance, Part 1, with the bonus section", "",
           "Generated by `python3 deck/build_deck.py` from `deck/slides.yaml` (Part 1) and `deck/slides_bonus.yaml`",
           "(the bonus section); the same text is in the pptx notes. Edit the yaml, not this file.",
           "Measured numbers for the bonus figures come from docs/bonus-results.md (deck/make_bonus_figures.py).", ""]
    for part, head in (("main", f"## Part 1: the workshop ({MAIN_MINUTES} minutes)"),
                       ("bonus", f"## Bonus: core principles of learning theory ({BONUS_MINUTES} minutes, its own clock)")):
        group = [s for s in slides if s["part"] == part]
        step = "Process step" if part == "main" else "Chapter"
        out += [head, "", f"| Slide | Title | Minutes | {step} | Notebook | Chip |", "|---|---|---|---|---|---|"]
        for s in group:
            a, b = s["minutes"]
            nb = f"▶ run section {s['run']}" if s.get("run") else ""
            title = (s["title"] if s["layout"] != "title" else "Title") + (" (optional)" if s.get("optional") else "")
            stage = {"none": "", "ahead": "(strip, none lit)"}.get(s["stage"], s["stage"])
            chip = plain(chip_text(s) or "")
            out.append(f"| {s['n']} | {title} | {a} to {b} | {stage} | {nb} | {chip} |")
        out.append("")
    for s in slides:
        a, b = s["minutes"]
        label = "Bonus minutes" if s["part"] == "bonus" else "Minutes"
        out.append(f"## {s['n']}. {s['title']}" + (" (optional)" if s.get("optional") else ""))
        out.append("")
        out.append(f"**{label} {a} to {b} ({b - a} min).**" +
                   (f" ▶ Notebook: run section {s['run']}, *{sections[s['run']]}*." if s.get("run") else ""))
        out.append("")
        if s.get("optional"):
            out.extend(["Optional slide: skip it when time is short; nothing later depends on it.", ""])
        out.append(s["notes"].strip())
        if s.get("cue"):
            out.extend(["", f"**Cue:** {s['cue']}"])
        if s.get("reading"):
            out.extend(["", f"{REF_LABEL}: {s['reading']}"])
        for c in s.get("reading_full", []):
            out.extend(["", f"Full citation: {c}"])
        out.append("")
    text = "\n".join(out)
    if DASHES.search(text):
        raise SystemExit("SPEAKER_NOTES.md would contain an em or en dash")
    OUT_NOTES.write_text(text, encoding="utf-8")


def build(force_figures=False):
    slides, sections, process, chapters, errs = load_spec()
    errs += self_check(slides, process, chapters)
    if errs:
        raise SystemExit("self-check failed:\n  " + "\n  ".join(errs))
    print("contrast against the WCAG floor:")
    for what, fg, bg, floor in CONTRAST_PAIRS:
        print(f"  {contrast(fg, bg):5.2f}:1  (floor {floor})  {what}: #{fg} on #{bg}")
    make_diagrams.make_all()
    ensure_result_figures(force_figures)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    # the default master uses an en dash as a level-2/4 bullet glyph; the house style bans the character
    for bu in prs.slide_masters[0].element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"):
        if DASHES.search(bu.get("char", "")):
            bu.set("char", "•")
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        paint_canvas(sl)
        LAYOUTS[s["layout"]](sl, s)
        if s["layout"] != "title":
            chrome(sl, s, sections, len(slides), process if s["part"] == "main" else chapters)
        sl.notes_slide.notes_text_frame.text = notes_text(s)
    prs.core_properties.title = "Machine Learning for Quantitative Finance, Part 1"
    prs.core_properties.author = "QuantSoc, Quantitative Finance Society"
    errs = check_built(prs, slides)
    if errs:
        raise SystemExit("post-build check failed:\n  " + "\n  ".join(errs))
    prs.save(OUT_PPTX)
    write_notes_md(slides, sections)
    print(f"wrote {OUT_PPTX.relative_to(ROOT)} ({len(slides)} slides: {MAIN_COUNT} in Part 1, {BONUS_COUNT} bonus)"
          f" and {OUT_NOTES.relative_to(ROOT)}")
    wanted = set()
    for s in slides:
        v = s.get("visual") or {}
        for it in [v] + v.get("items", []):
            if it.get("kind") == "result":
                wanted.add(it["file"])
    missing = sorted(f for f in wanted if not (FIG / f).exists())
    if missing:
        print("placeholders drawn for: " + ", ".join(missing))
    keys = sorted(set(re.findall(r"\{\{bonus:([A-Za-z0-9_]+)\}\}", OUT_NOTES.read_text(encoding="utf-8"))))
    if keys:
        print(f"{{{{bonus:KEY}}}} placeholders left in the notes ({len(keys)}): " + ", ".join(keys))
    ph = [c for c in make_qr.CODES if make_qr.is_placeholder(c[2])]
    if ph:
        print("QR placeholders still marked: " + ", ".join(c[1] for c in ph))


if __name__ == "__main__":
    build(force_figures="--figures" in sys.argv[1:])
